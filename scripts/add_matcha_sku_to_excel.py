import openpyxl
from copy import copy
import sys

sys.stdout.reconfigure(encoding='utf-8')

print("Opening Team_CamGiang_Report.xlsx...")
wb = openpyxl.load_workbook('Team_CamGiang_Report.xlsx')
ws = wb['Chi tiết SKU FamilyMart']

# Column 54 is BB. Column 55 is BC.
col_bc = 55
col_bb = 54

# 1. Header Row 4 (Category group)
ws.cell(row=4, column=col_bc, value="SẢN PHẨM MỚI PHÁT SINH / KHÁC")
src_cat = ws.cell(row=4, column=col_bb)
dst_cat = ws.cell(row=4, column=col_bc)
if src_cat.has_style:
    dst_cat.font = copy(src_cat.font)
    dst_cat.border = copy(src_cat.border)
    dst_cat.fill = copy(src_cat.fill)
    dst_cat.alignment = copy(src_cat.alignment)

# 2. Header Row 5 (SKU name & code)
sku_title = "Creamer đặc Ông Thọ vị Matcha tuýp 165g\n(01TM60)"
ws.cell(row=5, column=col_bc, value=sku_title)
src_h = ws.cell(row=5, column=col_bb)
dst_h = ws.cell(row=5, column=col_bc)
if src_h.has_style:
    dst_h.font = copy(src_h.font)
    dst_h.border = copy(src_h.border)
    dst_h.fill = copy(src_h.fill)
    dst_h.alignment = copy(src_h.alignment)

# Set column width
ws.column_dimensions['BC'].width = 22

# 3. Fill data for rows 6 to 40 and update col G formula
for r in range(6, 41):
    src_cell = ws.cell(row=r, column=col_bb)
    dst_cell = ws.cell(row=r, column=col_bc)
    
    # Store DC37 is at row 34
    addr_val = str(ws.cell(row=r, column=1).value or '')
    if 'DC37' in addr_val or 'Vietsing' in addr_val:
        dst_cell.value = 274800
        print(f"Row {r} (DC37): assigned 274,800 đ")
    else:
        dst_cell.value = 0
        
    if src_cell.has_style:
        dst_cell.font = copy(src_cell.font)
        dst_cell.border = copy(src_cell.border)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.number_format = copy(src_cell.number_format)
        dst_cell.alignment = copy(src_cell.alignment)
        
    # Update col G formula to include column BC
    ws.cell(row=r, column=7, value=f"=SUM(I{r}:BC{r})")

# 4. Total row (row 41)
total_row = ws.max_row
ws.cell(row=total_row, column=col_bc, value=f"=SUBTOTAL(9, BC6:BC{total_row-1})")
src_tot = ws.cell(row=total_row, column=col_bb)
dst_tot = ws.cell(row=total_row, column=col_bc)
if src_tot.has_style:
    dst_tot.font = copy(src_tot.font)
    dst_tot.border = copy(src_tot.border)
    dst_tot.fill = copy(src_tot.fill)
    dst_tot.number_format = copy(src_tot.number_format)
    dst_tot.alignment = copy(src_tot.alignment)

wb.save('Team_CamGiang_Report_Updated.xlsx')
print("Successfully saved Team_CamGiang_Report_Updated.xlsx with new SKU column BC (01TM60)!")

import shutil
try:
    shutil.copyfile('Team_CamGiang_Report_Updated.xlsx', 'Team_CamGiang_Report.xlsx')
    print("Replaced Team_CamGiang_Report.xlsx successfully!")
    import os
    os.remove('Team_CamGiang_Report_Updated.xlsx')
except Exception as e:
    print("Note: Excel file is currently open in Excel application. Team_CamGiang_Report_Updated.xlsx is ready to be swapped once Excel is closed.")

